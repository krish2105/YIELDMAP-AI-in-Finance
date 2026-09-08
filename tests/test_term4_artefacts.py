"""The report and the deck must not be able to launder generated figures.

CLAUDE.md's rule is that generated data cannot reach a Term 4 artefact, and
`scripts/guard_synthetic.py` is the mechanism. These documents are the artefacts it exists to
stop, so the interesting tests are the refusals rather than the happy path.

The stamping matters as much as the refusal. A draft that generates but labels nothing is worse
than no draft: it looks finished.
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx", reason="artefact generation needs the `docs` extra")
pytest.importorskip("pptx", reason="artefact generation needs the `docs` extra")

from scripts import build_deck, build_report  # noqa: E402
from scripts.term4.artefacts import collect  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "AI217_YIELDMAP_report.docx"
DECK = ROOT / "docs" / "AI217_YIELDMAP_deck.pptx"


class TestTheGuardIsLoadBearing:
    """Generation is refused while the guard blocks, unless the caller says --draft."""

    @pytest.mark.parametrize("module", [build_report, build_deck])
    def test_it_refuses_without_draft_while_the_guard_blocks(self, module, tmp_path, capsys):
        out = tmp_path / "refused.bin"

        code = module.main(["--out", str(out)])

        assert code == 1, "generated a graded artefact from data the guard rejects"
        assert not out.exists(), "wrote the file anyway"
        assert "provenance guard is blocking" in capsys.readouterr().err

    @pytest.mark.parametrize(("module", "suffix"), [(build_report, ".docx"), (build_deck, ".pptx")])
    def test_draft_generates_a_readable_file(self, module, suffix, tmp_path):
        out = tmp_path / f"draft{suffix}"

        assert module.main(["--draft", "--out", str(out)]) == 0
        assert out.exists() and out.stat().st_size > 10_000
        with zipfile.ZipFile(out) as archive:
            assert archive.testzip() is None


class TestNothingIsRestatedFromMemory:
    def test_every_row_names_the_file_it_came_from(self):
        bundle = collect()
        rows = (
            bundle.rows_models
            + bundle.rows_retrieval
            + bundle.rows_engineering
            + bundle.rows_governance
        )
        assert rows, "the loader found no results at all"
        for row in rows:
            assert row.source.endswith(".json"), f"{row.label!r} cites no artefact"
            assert (ROOT / "docs" / "results" / row.source).exists() or (
                ROOT / "docs" / "results" / "synthetic" / row.source
            ).exists(), f"{row.label!r} cites {row.source}, which is not in docs/results/"

    def test_provenance_is_read_from_the_file_not_guessed(self):
        bundle = collect()
        for row in bundle.rows_models:
            body = json.loads((ROOT / "docs" / "results" / "synthetic" / row.source).read_text())
            assert row.provenance == body["provenance"]

    def test_a_missing_model_becomes_a_row_that_says_so_rather_than_a_number(self):
        from scripts.term4.artefacts import MISSING, num, pct

        assert num(None) == MISSING
        assert pct(None) == MISSING


class TestTheDraftIsStamped:
    """Three places, because one is a place a reader can skip."""

    def test_the_report_carries_the_banner_the_column_and_the_section(self):
        from docx import Document

        doc = Document(REPORT)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "DRAFT" in "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells), (
            "no draft banner on the cover"
        )
        assert "How provenance is enforced" in text, "no section explaining the guard"

        model_table = next(
            t for t in doc.tables if any("hedonic.json" in c.text for r in t.rows for c in r.cells)
        )
        headers = [c.text for c in model_table.rows[0].cells]
        assert "Data" in headers, "the model table has no provenance column"
        assert any("stand-in" in c.text for r in model_table.rows[1:] for c in r.cells), (
            "no row is marked as a stand-in"
        )

    def test_the_deck_names_the_stand_in_on_the_opening_slide(self):
        from pptx import Presentation

        prs = Presentation(DECK)
        first = " ".join(
            s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame
        ).upper()
        assert "STAND-IN" in first

    def test_every_slide_has_speaker_notes(self):
        from pptx import Presentation

        for i, slide in enumerate(Presentation(DECK).slides, start=1):
            assert slide.has_notes_slide, f"slide {i} has no notes"
            assert len(slide.notes_slide.notes_text_frame.text) > 80, f"slide {i} notes are thin"


class TestTheDocumentsFitTheirPages:
    """The defect a renderer would show, caught by arithmetic instead."""

    def test_no_table_is_wider_than_the_text_column(self):
        from docx import Document

        doc = Document(REPORT)
        section = doc.sections[0]
        text_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
        for i, table in enumerate(doc.tables):
            width = sum(c.width.cm for c in table.rows[0].cells if c.width)
            assert width <= text_cm + 0.01, (
                f"table {i} is {width:.2f}cm in a {text_cm:.2f}cm column"
            )

    def test_no_shape_falls_off_a_slide(self):
        from pptx import Presentation

        prs = Presentation(DECK)
        for i, slide in enumerate(prs.slides, start=1):
            for shape in slide.shapes:
                if shape.left is None:
                    continue
                assert shape.left >= 0 and shape.top >= 0, f"slide {i}: shape starts off-slide"
                assert shape.left + shape.width <= prs.slide_width, f"slide {i}: shape overruns"
                assert shape.top + shape.height <= prs.slide_height, f"slide {i}: shape overruns"
